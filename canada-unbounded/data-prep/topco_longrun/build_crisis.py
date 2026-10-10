"""Build tidy crisis-comparison CSV (GFC 2008-09 and COVID-19) for G7 + OECD.
Usage: python3 -I build_crisis.py <raw_dir> <out_dir>
"""
import csv, json, sys, os, collections, datetime

RAW, OUT = sys.argv[1], sys.argv[2]
G7 = ["CAN", "USA", "GBR", "DEU", "FRA", "ITA", "JPN"]
OECD = "AUS AUT BEL CAN CHL COL CRI CZE DNK EST FIN FRA DEU GRC HUN ISL IRL ISR ITA JPN KOR LVA LTU LUX MEX NLD NZL NOR POL PRT SVK SVN ESP SWE CHE TUR GBR USA".split()
NAMES = {}
rows = []

SRC = {
    "wb_gdp": ("World Bank WDI NY.GDP.MKTP.KD.ZG (accessed 2026-10-10)", "https://api.worldbank.org/v2/country/all/indicator/NY.GDP.MKTP.KD.ZG"),
    "wb_gdpk": ("World Bank WDI NY.GDP.MKTP.KD, constant 2015 US$ (accessed 2026-10-10)", "https://api.worldbank.org/v2/country/all/indicator/NY.GDP.MKTP.KD"),
    "wb_pop": ("World Bank WDI SP.POP.TOTL", "https://api.worldbank.org/v2/country/all/indicator/SP.POP.TOTL"),
    "imf_gdp": ("IMF World Economic Outlook April 2026 via DataMapper NGDP_RPCH", "https://www.imf.org/external/datamapper/api/v1/NGDP_RPCH"),
    "imf_lur": ("IMF World Economic Outlook April 2026 via DataMapper LUR", "https://www.imf.org/external/datamapper/api/v1/LUR"),
    "oecd_qna": ("OECD Quarterly National Accounts, real GDP chain-linked volume, SA (DF_QNA, accessed 2026-10-10)", "https://sdmx.oecd.org/public/rest/data/OECD.SDD.NAD,DSD_NAMAIN1@DF_QNA,1.1/"),
    "oecd_une": ("OECD monthly harmonised unemployment rate, SA, 15+ (DF_IALFS_UNE_M)", "https://sdmx.oecd.org/public/rest/data/OECD.SDD.TPS,DSD_LFS@DF_IALFS_UNE_M,1.0/"),
    "oecd_xm": ("OECD Excess mortality by week (DF_MORTALITY_EXCESS; excess vs 2015-19 average, not age-adjusted)", "https://sdmx.oecd.org/public/rest/data/OECD.ELS.HD,DSD_HEALTH_MORTALITY@DF_MORTALITY_EXCESS,1.0/"),
    "who_xm": ("WHO, Estimates of excess mortality associated with COVID-19 (as of 5 April 2023), modelled", "https://www.who.int/data/sets/global-excess-deaths-associated-with-covid-19-modelled-estimates"),
    "owid_hmd": ("Our World in Data excess mortality (HMD/WMD data, projected baseline; Karlinsky & Kobak method) - aggregator", "https://github.com/owid/covid-19-data/tree/master/public/data/excess_mortality"),
    "econ": ("The Economist excess-deaths model via OWID - media model estimate", "https://github.com/owid/covid-19-data/blob/master/public/data/excess_mortality/excess_mortality_economist_estimates.csv"),
    "owid_vax": ("Our World in Data COVID-19 vaccinations (official national data compiled by OWID)", "https://github.com/owid/covid-19-data/blob/master/public/data/vaccinations/vaccinations.csv"),
    "imf_fm": ("IMF Fiscal Monitor Database of Country Fiscal Measures in Response to COVID-19, Table 1 (as of 27 Sep 2021), % of 2020 GDP", "https://meetings.imf.org/-/media/Files/Topics/COVID/FM-Database/October-2021/oct-2021-country-fiscal-measures-database-publication.ashx"),
    "lv": ("Laeven & Valencia (2018), Systemic Banking Crises Revisited, IMF WP/18/206, Appendix Tables 1-2", "https://www.imf.org/en/Publications/WP/Issues/2018/09/14/Systemic-Banking-Crises-Revisited-46232"),
}

def add(country, period, indicator, value, unit, src, note=""):
    if value is None:
        return
    name, url = SRC[src]
    rows.append(dict(country=country, period=period, indicator=indicator,
                     value=round(value, 3) if isinstance(value, float) else value,
                     unit=unit, source=name, source_url=url, note=note))

# ---------- World Bank ----------
def wb(fn):
    d = json.load(open(os.path.join(RAW, "wb", fn)))
    out = collections.defaultdict(dict)
    for r in d[1]:
        NAMES[r["countryiso3code"]] = r["country"]["value"]
        if r["value"] is not None:
            out[r["countryiso3code"]][int(r["date"])] = r["value"]
    return out
wbg = wb("oecd_NY.GDP.MKTP.KD.ZG.json")
wbk = wb("oecd_NY.GDP.MKTP.KD.json")
pop = wb("oecd_SP.POP.TOTL.json")
for c in OECD:
    for y in (2008, 2009, 2010, 2020, 2021):
        add(c, str(y), "real_gdp_growth", wbg[c].get(y), "% y/y", "wb_gdp")
    k = wbk[c]
    if 2007 in k and 2009 in k:
        add(c, "2007-2009", "real_gdp_change_2009_vs_2007", (k[2009] / k[2007] - 1) * 100, "%", "wb_gdpk")
    if 2008 in k:
        pk = max(k.get(2007, 0), k[2008])
        rec = next((y for y in range(2009, 2026) if y in k and k[y] >= pk), None)
        add(c, "GFC", "first_year_annual_gdp_above_2007_08_peak", rec, "year", "wb_gdpk")
    if 2019 in k:
        rec = next((y for y in range(2020, 2026) if y in k and k[y] >= k[2019]), None)
        add(c, "COVID", "first_year_annual_gdp_above_2019", rec, "year", "wb_gdpk")
        if 2021 in k:
            add(c, "2019-2021", "real_gdp_2021_vs_2019", (k[2021] / k[2019] - 1) * 100, "%", "wb_gdpk")
        if 2023 in k:
            add(c, "2019-2023", "real_gdp_2023_vs_2019", (k[2023] / k[2019] - 1) * 100, "%", "wb_gdpk")

# ---------- IMF WEO ----------
imf = json.load(open(os.path.join(RAW, "imf", "weo_ngdp_rpch_all.json")))["values"]["NGDP_RPCH"]
lur = json.load(open(os.path.join(RAW, "imf", "weo_lur.json")))["values"]["LUR"]
for c in OECD:
    for y in ("2008", "2009", "2020", "2021"):
        v = imf.get(c, {}).get(y)
        add(c, y, "real_gdp_growth_weo", v, "% y/y", "imf_gdp")
    for y in ("2007", "2009", "2010", "2019", "2020", "2021"):
        add(c, y, "unemployment_rate_annual_weo", lur.get(c, {}).get(y), "%", "imf_lur")

# ---------- OECD quarterly GDP ----------
q = collections.defaultdict(dict)
for r in csv.DictReader(open(os.path.join(RAW, "oecd", "qgdp_oecd.csv"))):
    if r["OBS_VALUE"]:
        q[r["REF_AREA"]][r["TIME_PERIOD"]] = float(r["OBS_VALUE"])
def qrange(a, b, keys):
    return [k for k in sorted(keys) if a <= k <= b]
for c in OECD:
    s = q.get(c)
    if not s:
        continue
    # GFC: peak 2007Q1-2009Q2, trough after peak through 2010Q4
    pre = qrange("2007-Q1", "2009-Q2", s)
    if pre:
        pk = max(pre, key=lambda k: s[k])
        post = qrange(pk, "2010-Q4", s)
        tr = min(post, key=lambda k: s[k])
        add(c, "GFC", "gdp_peak_quarter", pk, "quarter", "oecd_qna")
        add(c, "GFC", "gdp_trough_quarter", tr, "quarter", "oecd_qna")
        add(c, "GFC", "gdp_peak_to_trough", (s[tr] / s[pk] - 1) * 100, "%", "oecd_qna")
        later = [k for k in sorted(s) if k > tr and s[k] >= s[pk]]
        if later:
            rq = later[0]
            n = (int(rq[:4]) - int(pk[:4])) * 4 + int(rq[-1]) - int(pk[-1])
            add(c, "GFC", "gdp_quarter_regained_precrisis_peak", rq, "quarter", "oecd_qna")
            add(c, "GFC", "gdp_quarters_peak_to_recovery", n, "quarters", "oecd_qna")
    # COVID: base 2019Q4
    if "2019-Q4" in s:
        b = s["2019-Q4"]
        tr = min(qrange("2020-Q1", "2020-Q4", s), key=lambda k: s[k])
        add(c, "COVID", "gdp_trough_quarter", tr, "quarter", "oecd_qna")
        add(c, "COVID", "gdp_trough_vs_2019Q4", (s[tr] / b - 1) * 100, "%", "oecd_qna")
        later = [k for k in sorted(s) if k > tr and s[k] >= b]
        if later:
            add(c, "COVID", "gdp_quarter_regained_2019Q4", later[0], "quarter", "oecd_qna")
            rq = later[0]
            n = (int(rq[:4]) - 2019) * 4 + int(rq[-1]) - 4
            add(c, "COVID", "gdp_quarters_2019Q4_to_recovery", n, "quarters", "oecd_qna")
        if "2021-Q4" in s:
            add(c, "COVID", "gdp_2021Q4_vs_2019Q4", (s["2021-Q4"] / b - 1) * 100, "%", "oecd_qna")
        if "2023-Q4" in s:
            add(c, "COVID", "gdp_2023Q4_vs_2019Q4", (s["2023-Q4"] / b - 1) * 100, "%", "oecd_qna")

# ---------- OECD monthly unemployment ----------
u = collections.defaultdict(dict)
for r in csv.DictReader(open(os.path.join(RAW, "oecd", "unemp_m.csv"))):
    if r["OBS_VALUE"]:
        u[r["REF_AREA"]][r["TIME_PERIOD"]] = float(r["OBS_VALUE"])
for c in OECD:
    s = u.get(c)
    if not s:
        continue
    pre = [s[k] for k in s if k.startswith("2007")]
    win = qrange("2008-01", "2012-12", s)
    if pre and win:
        pk = max(win, key=lambda k: s[k])
        add(c, "GFC", "unemployment_2007_avg", sum(pre) / len(pre), "%", "oecd_une")
        add(c, "GFC", "unemployment_peak_2008_2012", s[pk], "%", "oecd_une", note=f"peak month {pk}")
        add(c, "GFC", "unemployment_rise_pp_vs_2007_avg", s[pk] - sum(pre) / len(pre), "pp", "oecd_une")
    if "2020-02" in s:
        win = qrange("2020-03", "2021-12", s)
        pk = max(win, key=lambda k: s[k])
        add(c, "COVID", "unemployment_feb2020", s["2020-02"], "%", "oecd_une")
        add(c, "COVID", "unemployment_peak_2020_2021", s[pk], "%", "oecd_une", note=f"peak month {pk}")
        add(c, "COVID", "unemployment_rise_pp_vs_feb2020", s[pk] - s["2020-02"], "pp", "oecd_une")

# ---------- Excess mortality: OECD weekly ----------
xm = collections.defaultdict(lambda: collections.defaultdict(float))
for r in csv.DictReader(open(os.path.join(RAW, "oecd", "excess_week.csv"))):
    if r["UNIT_MEASURE"] == "DT" and r["AGE"] == "_T" and r["SEX"] == "_T" and r["OBS_VALUE"]:
        y = r["TIME_PERIOD"][:4]
        xm[r["REF_AREA"]][y] += float(r["OBS_VALUE"])
for c in OECD:
    if c not in xm:
        continue
    p = pop[c].get(2020)
    for per, ys in (("2020-2021", ("2020", "2021")), ("2020-2022", ("2020", "2021", "2022")), ("2020-2023", ("2020", "2021", "2022", "2023"))):
        if all(y in xm[c] for y in ys) and p:
            tot = sum(xm[c][y] for y in ys)
            add(c, per, "excess_deaths_per_million_oecd", tot / p * 1e6, "per million (2020 pop.)", "oecd_xm",
                note="sum of weekly excess deaths vs 2015-19 average; ISO weeks; not age-standardised")

# ---------- WHO ----------
import openpyxl
wbx = openpyxl.load_workbook(os.path.join(RAW, "who", "gem", "2023-05-19_covid-19_gem", "WHO_COVID_Excess_Deaths_EstimatesByCountry.xlsx"), read_only=True)
exc = {}
for r in wbx["Deaths by year"].iter_rows(values_only=True):
    if r and r[1] in OECD and r[2] == "2020-2021":
        exc[r[1]] = (r[3], r[4], r[5])
expd = collections.defaultdict(float)
for r in wbx["Deaths by year and month"].iter_rows(values_only=True):
    if r and r[1] in OECD and r[2] in (2020, 2021) and isinstance(r[5], (int, float)):
        expd[r[1]] += r[5]
for c, (m, lo, hi) in exc.items():
    p = pop[c].get(2020)
    if p:
        add(c, "2020-2021", "excess_deaths_per_100k_who", m / p * 1e5, "per 100k (WB 2020 pop.)", "who_xm",
            note=f"95% CI per 100k {lo / p * 1e5:.1f} to {hi / p * 1e5:.1f}")
    if expd[c]:
        add(c, "2020-2021", "excess_deaths_pct_of_expected_who", m / expd[c] * 100, "% (P-score)", "who_xm")

# ---------- OWID HMD/WMD (projected baseline) ----------
iso = {"Canada": "CAN", "United States": "USA", "United Kingdom": "GBR", "Germany": "DEU", "France": "FRA", "Italy": "ITA", "Japan": "JPN"}
last = {}
for r in csv.DictReader(open(os.path.join(RAW, "owid", "excess_mortality.csv"))):
    c = iso.get(r["location"])
    if c and r["cum_excess_per_million_proj_all_ages"] and r["date"] <= "2022-12-31":
        if c not in last or r["date"] > last[c][0]:
            last[c] = (r["date"], float(r["cum_excess_per_million_proj_all_ages"]), r["cum_p_proj_all_ages"])
for c, (d, v, pp) in last.items():
    add(c, "2020-2022", "excess_deaths_per_million_owid_projected_baseline", v, "per million", "owid_hmd", note=f"cumulative to {d}")
    if pp:
        add(c, "2020-2022", "excess_deaths_pct_owid_projected_baseline", float(pp), "% (cumulative P-score)", "owid_hmd", note=f"cumulative to {d}")

# ---------- Economist model ----------
ec = {}
for r in csv.DictReader(open(os.path.join(RAW, "owid", "excess_mortality_economist_estimates.csv"))):
    c = iso.get(r["country"])
    if c and r["date"] <= "2022-12-31" and r["cumulative_estimated_daily_excess_deaths_per_100k"]:
        if c not in ec or r["date"] > ec[c][0]:
            ec[c] = (r["date"], float(r["cumulative_estimated_daily_excess_deaths_per_100k"]))
for c, (d, v) in ec.items():
    add(c, "2020-2022", "excess_deaths_per_100k_economist_model", v, "per 100k", "econ", note=f"cumulative to {d}")

# ---------- Vaccination ----------
vx = collections.defaultdict(dict)
for r in csv.DictReader(open(os.path.join(RAW, "owid", "vaccinations.csv"))):
    if r["iso_code"] in OECD:
        vx[r["iso_code"]][r["date"]] = r
for c in OECD:
    for tgt, lab in (("2021-06-30", "2021-06-30"), ("2021-12-31", "2021-12-31"), ("2022-12-31", "2022-12-31")):
        ds = [d for d in vx[c] if d <= tgt and vx[c][d]["people_fully_vaccinated_per_hundred"]]
        if ds:
            d = max(ds)
            add(c, lab, "people_fully_vaccinated_per_hundred", float(vx[c][d]["people_fully_vaccinated_per_hundred"]), "% of population",
                "owid_vax", note=f"latest obs {d}")
        ds = [d for d in vx[c] if d <= tgt and vx[c][d]["people_vaccinated_per_hundred"]]
        if ds:
            d = max(ds)
            add(c, lab, "people_at_least_one_dose_per_hundred", float(vx[c][d]["people_vaccinated_per_hundred"]), "% of population",
                "owid_vax", note=f"latest obs {d}")

# ---------- IMF fiscal measures (transcribed from Table 1 text; see notes) ----------
fm = {  # above-the-line subtotal, health, non-health, liquidity subtotal (% of 2020 GDP)
    "AUS": (18.4, 1.0, 17.4, 1.8), "CAN": (15.9, 2.8, 13.1, 4.0), "FRA": (9.6, 1.5, 8.2, 15.2),
    "DEU": (15.3, 1.8, 13.6, 27.8), "ITA": (10.9, 1.2, 9.7, 35.3), "JPN": (16.7, 2.1, 14.6, 28.3),
    "GBR": (19.3, 4.8, 14.4, 16.7), "USA": (25.5, 3.3, 22.2, 2.4)}
for c, (a, h, nh, l) in fm.items():
    add(c, "2020-2021+", "fiscal_above_the_line_pct_gdp", a, "% of 2020 GDP", "imf_fm", note="additional spending + forgone revenue")
    add(c, "2020-2021+", "fiscal_above_the_line_health_pct_gdp", h, "% of 2020 GDP", "imf_fm")
    add(c, "2020-2021+", "fiscal_above_the_line_nonhealth_pct_gdp", nh, "% of 2020 GDP", "imf_fm")
    add(c, "2020-2021+", "fiscal_liquidity_support_pct_gdp", l, "% of 2020 GDP", "imf_fm", note="equity, loans, asset purchases, guarantees, quasi-fiscal")

# ---------- Laeven-Valencia systemic banking crisis ----------
lv = {"CAN": ("none", None, None), "JPN": ("none (2007-09)", None, None), "USA": ("2007", 30.0, 4.5), "GBR": ("2007", 25.3, 8.8),
      "DEU": ("2008", 12.3, 2.7), "FRA": ("2008 (borderline)", 23.3, 1.3), "ITA": ("2008", 32.2, 0.7)}
for c, (st, ol, fc) in lv.items():
    add(c, "GFC", "systemic_banking_crisis_start", st, "year", "lv")
    add(c, "GFC", "banking_crisis_output_loss_pct_trend_gdp", ol, "% of trend GDP (T..T+3)", "lv")
    add(c, "GFC", "banking_crisis_gross_fiscal_cost_pct_gdp", fc, "% of GDP", "lv", note="outlays for financial-sector restructuring")

# ---------- write tidy ----------
os.makedirs(OUT, exist_ok=True)
with open(os.path.join(OUT, "crisis_comparison_long.csv"), "w", newline="") as f:
    w = csv.DictWriter(f, fieldnames=["country", "country_name", "g7", "period", "indicator", "value", "unit", "source", "source_url", "note"])
    w.writeheader()
    for r in rows:
        r["country_name"] = NAMES.get(r["country"], r["country"])
        r["g7"] = r["country"] in G7
        w.writerow(r)

# ---------- ranks among G7 and OECD ----------
lower_better = {"real_gdp_growth": False, "real_gdp_change_2009_vs_2007": False, "gdp_peak_to_trough": False,
                "gdp_quarters_peak_to_recovery": True, "unemployment_rise_pp_vs_2007_avg": True,
                "gdp_trough_vs_2019Q4": False, "gdp_quarters_2019Q4_to_recovery": True, "gdp_2021Q4_vs_2019Q4": False,
                "gdp_2023Q4_vs_2019Q4": False, "real_gdp_2021_vs_2019": False, "real_gdp_2023_vs_2019": False,
                "unemployment_rise_pp_vs_feb2020": True, "unemployment_peak_2020_2021": True, "unemployment_peak_2008_2012": True,
                "excess_deaths_per_million_oecd": True, "excess_deaths_per_100k_who": True, "excess_deaths_pct_of_expected_who": True,
                "excess_deaths_per_million_owid_projected_baseline": True, "excess_deaths_per_100k_economist_model": True,
                "people_fully_vaccinated_per_hundred": False, "real_gdp_growth_weo": False}
grp = collections.defaultdict(dict)
for r in rows:
    if r["indicator"] in lower_better and isinstance(r["value"], (int, float)):
        grp[(r["indicator"], r["period"])][r["country"]] = r["value"]
with open(os.path.join(OUT, "crisis_canada_ranks.csv"), "w", newline="") as f:
    w = csv.writer(f)
    w.writerow(["indicator", "period", "better_is", "canada_value", "canada_rank_g7", "n_g7", "best_g7", "best_g7_value", "canada_rank_oecd", "n_oecd", "best_oecd", "best_oecd_value"])
    for (ind, per), d in sorted(grp.items()):
        if "CAN" not in d:
            continue
        lb = lower_better[ind]
        def rank(cs):
            cs = [c for c in cs if c in d]
            srt = sorted(cs, key=lambda c: d[c], reverse=not lb)
            return srt.index("CAN") + 1, len(srt), srt[0], d[srt[0]]
        rg = rank(G7); ro = rank(OECD)
        w.writerow([ind, per, "lower" if lb else "higher", d["CAN"], *rg, *ro])
print(len(rows), "rows")
